"""Build the Portuguese project report from verified project artifacts."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "reports" / "relatorio_deteccao_malware.docx"
NAVY = "17324D"
BLUE = "247BA0"
PALE = "EAF2F6"
GREY = "53616D"


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, value: str, *, bold: bool = False, color: str = "263746") -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    run = paragraph.add_run(str(value))
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(8.4)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Shading Accent 1"
    for index, heading in enumerate(headers):
        set_cell_text(table.rows[0].cells[index], heading, bold=True, color="FFFFFF")
        shade(table.rows[0].cells[index], NAVY)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
            if len(table.rows) % 2 == 0:
                shade(cells[index], "F3F7F9")
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                row.cells[index].width = Inches(width)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_paragraph(document: Document, text: str, *, lead: str | None = None) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.05
    if lead and text.startswith(lead):
        first = paragraph.add_run(lead)
        first.bold = True
        paragraph.add_run(text[len(lead):])
    else:
        paragraph.add_run(text)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.0
        paragraph.add_run(item)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    heading = document.add_heading(text, level=level)
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.space_before = Pt(7 if level == 1 else 4)
    heading.paragraph_format.space_after = Pt(4)


def add_figure(document: Document, relative_path: str, caption: str, width: float = 6.0) -> None:
    path = ROOT / relative_path
    if path.is_file():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(1)
        paragraph.add_run().add_picture(str(path), width=Inches(width))
        caption_paragraph = document.add_paragraph()
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_paragraph.paragraph_format.space_after = Pt(4)
        run = caption_paragraph.add_run(caption)
        run.italic = True
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(GREY)
    else:
        add_paragraph(document, f"Figura não incluída: artefacto não encontrado ({relative_path}).")


def add_page(document: Document, title: str, contents: list[tuple[str, object]]) -> None:
    document.add_page_break()
    add_heading(document, title, 1)
    for kind, value in contents:
        if kind == "p":
            add_paragraph(document, str(value))
        elif kind == "bullets":
            add_bullets(document, value)  # type: ignore[arg-type]
        elif kind == "table":
            headers, rows, *widths = value  # type: ignore[misc]
            add_table(document, headers, rows, widths[0] if widths else None)
        elif kind == "figure":
            relative_path, caption, *widths = value  # type: ignore[misc]
            add_figure(document, relative_path, caption, widths[0] if widths else 6.0)
        elif kind == "subheading":
            add_heading(document, str(value), 2)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def build_report() -> Path:
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.65)
    section.bottom_margin = Cm(1.55)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.header_distance = Cm(0.7)
    section.footer_distance = Cm(0.65)

    styles = document.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(9.3)
    styles["Normal"].font.color.rgb = RGBColor.from_string("263746")
    for name, size, color in (("Title", 28, NAVY), ("Heading 1", 17, NAVY), ("Heading 2", 11, BLUE)):
        styles[name].font.name = "Aptos Display"
        styles[name].font.size = Pt(size)
        styles[name].font.bold = True
        styles[name].font.color.rgb = RGBColor.from_string(color)
    styles["List Bullet"].font.name = "Aptos"
    styles["List Bullet"].font.size = Pt(9)

    header = section.header.paragraphs[0]
    header.text = "IAAC  |  Relatório de projeto"
    header.style = styles["Normal"]
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string(GREY)
    footer = section.footer.paragraphs[0]
    add_page_number(footer)

    # Capa
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(90)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("IAAC")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(30)
    title.paragraph_format.space_after = Pt(14)
    run = title.add_run("Deteção de malware\ncom aprendizagem automática")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(27)
    run.font.color.rgb = RGBColor.from_string(NAVY)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Relatório técnico e de enquadramento do projeto")
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string(GREY)
    document.add_paragraph()
    add_table(document, ["Identificação académica", "Preencher antes da entrega"], [
        ["Estudante(s)", "[A preencher pela equipa]"],
        ["Unidade curricular / curso", "[A preencher]"],
        ["Docente", "[A preencher]"],
        ["Instituição", "ISTEC"],
        ["Data", "7 de outubro de 2026"],
    ], [2.4, 3.6])
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(38)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Protótipo académico; não substitui análise profissional nem constitui um produto de segurança validado.")
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(GREY)

    pages: list[tuple[str, list[tuple[str, object]]]] = [
        ("Resumo executivo", [
            ("p", "Este projeto constrói um fluxo local para explorar dados de ficheiros Windows e classificar amostras em Benign ou Malware. A abordagem combina preparação reprodutível, comparação de classificadores, avaliação em dados reservados e uma interface de utilização. O objetivo é apoiar a triagem humana; não é afirmar que o sistema deteta todas as formas de ransomware."),
            ("p", "Na versão avaliada, a limpeza reduziu 21.752 registos para 13.886, após deduplicação por MD5 e remoção documentada de 611 registos com etiquetas incompatíveis. O conjunto final contém 72 características e as classes têm dimensão semelhante: 7.168 Benign e 6.718 Malware."),
            ("p", "O Random Forest foi escolhido por recall de Malware na validação e ajustado por validação cruzada estratificada. No conjunto de teste reservado (2.778 linhas), obteve recall 0,9970, precisão 0,9875, F1 0,9922, ROC-AUC 0,9990 e PR-AUC 0,9989. Houve 4 falsos negativos e 17 falsos positivos. Estes números descrevem este conjunto e este protocolo; não provam desempenho em dados atuais, famílias novas ou ambientes reais."),
            ("p", "A implementação extrai metadados PE estaticamente de executáveis e ZIPs, dispõe de uma aplicação de ambiente de trabalho, de uma API local e de uma integração opcional com Telegram. A extração não executa o ficheiro. As características comportamentais não são obtidas por esta análise estática e podem não existir nos ficheiros enviados."),
            ("p", "As prioridades seguintes são validar por hash e por origem temporal, testar dados externos, avaliar o comportamento quando faltam campos e fechar controlos operacionais antes de qualquer utilização em rede. A declaração de apoio de IA e os limites são apresentados no final.")
        ]),
        ("1. Enquadramento e âmbito", [
            ("p", "A quantidade de ficheiros que uma equipa de segurança pode ter de rever torna útil uma primeira ordenação automática. O projeto explora se características estruturais de executáveis PE, acompanhadas por variáveis do dataset, ajudam a separar ficheiros rotulados Benign de Malware."),
            ("subheading", "Pergunta do projeto"),
            ("p", "Com as características disponíveis, consegue-se priorizar amostras rotuladas como Malware com recall elevado, sem tornar excessivo o número de alertas falsos? A resposta depende do conjunto, da qualidade das etiquetas, da origem dos dados e do método de divisão."),
            ("p", "O alvo implementado é binário: Class ∈ {Benign, Malware}. Embora o nome do projeto refira ransomware, o conjunto contém também RAT, Stealer e Trojan. Logo, os resultados não são uma avaliação dedicada a ransomware. A análise multiclasse por categoria ou família é trabalho futuro, não uma capacidade demonstrada pelo resultado binário."),
            ("table", (["Incluído", "Fora do âmbito demonstrado"], [
                ["EDA, preparação, comparação e avaliação supervisionada", "Deteção garantida de todas as variantes ou famílias"],
                ["Extração estática de características PE", "Execução em sandbox / análise dinâmica"],
                ["API e interface para protótipo local", "Serviço de produção, monitorização operacional e resposta automática"],
                ["Triagem assistida por classificação binária", "Identificação inequívoca de ransomware ou prova de segurança"],
            ]))
        ]),
        ("2. Canvas: tarefa de previsão e decisão", [
            ("p", "A estrutura de enquadramento segue as dez áreas do Machine Learning Canvas: da decisão que se pretende apoiar até à monitorização. Aqui, cada área é preenchida com factos do projeto ou assinalada como hipótese ainda por confirmar."),
            ("subheading", "Tarefa de previsão"),
            ("table", (["Elemento", "Definição no projeto"], [
                ["Unidade", "Um registo de ficheiro representado por características tabulares."],
                ["Entrada", "Características do dataset ou atributos PE extraídos localmente."],
                ["Saída", "Etiqueta binária Benign / Malware e, quando disponível, score do modelo."],
                ["Horizonte", "Previsão para a amostra analisada; não há previsão temporal."],
                ["Utilizador", "Analista / utilizador que revê o alerta; a decisão final é humana."],
            ])),
            ("subheading", "Decisão que a previsão apoia"),
            ("p", "A saída pode sugerir a ordem de triagem: rever primeiro uma amostra classificada como Malware. A previsão não deve apagar, bloquear, quarentenar nem aprovar automaticamente um ficheiro. Uma decisão de produção precisaria de política de risco, limiar calibrado, fallback para análise especializada e registo auditável; esses controlos não estão demonstrados pelo benchmark.")
        ]),
        ("3. Canvas: proposta de valor e impacto", [
            ("p", "A proposta de valor é reduzir o esforço de triagem inicial e tornar a avaliação repetível. A utilidade prática não foi medida com analistas, tempos observados ou comparação controlada entre fluxo manual e fluxo assistido."),
            ("p", "O valor técnico mensurável nesta fase é a capacidade de treinar e avaliar um classificador em dados tabulares, apresentar métricas e explorar entradas por uma interface local. O ganho operacional ou financeiro permanece por quantificar."),
            ("subheading", "Simulação de impacto: o que falta medir"),
            ("p", "O Canvas de referência contém exemplos ilustrativos de custos de erro. Esses valores não são dados do projeto e não são usados como estimativas. Para calcular impacto, a equipa teria de observar volume diário, prevalência de malware, custo por falso positivo, custo por falso negativo, tempo médio de revisão e taxa de resolução após investigação."),
            ("table", (["Pergunta operacional", "Medição necessária"], [
                ["Menos tempo por amostra?", "Tempo de triagem com e sem apoio, numa amostra representativa."],
                ["Custo de alerta falso?", "Horas de analista e impacto de interrupções por alerta."],
                ["Custo de ameaça não detetada?", "Estimativa acordada com responsável de segurança e contexto de uso."],
                ["Benefício líquido?", "Volume, taxas reais, custo de operação e manutenção do modelo."],
            ]))
        ]),
        ("4. Canvas: recolha de dados e fontes", [
            ("p", "O principal artefacto de dados é ransom.csv, com 21.752 linhas e 77 colunas antes da limpeza. A documentação do projeto identifica o conjunto como associado ao Ransomware Dataset 2024 / Kaggle; o ficheiro local não inclui uma cadeia de proveniência independente que permita confirmar origem, licenças, datas de recolha ou processo de rotulagem."),
            ("p", "As colunas incluem identificadores (MD5/SHA-1), informação categórica e características PE, além de campos comportamentais (por exemplo, rede, registry, processos, ficheiros e APIs). A presença de um campo no dataset não significa que o extrator local o consiga observar para um ficheiro isolado."),
            ("p", "O estado inicial descrito na documentação assinalava 7.849 MD5 repetidos. O processamento analítico retirou duplicados e conflitos e conservou uma linha por MD5. A contagem de 611 linhas removidas por etiquetas conflituosas é explícita no relatório de EDA; esta escolha reduz ruído, mas também retira casos que mereciam resolução por proveniência."),
            ("subheading", "Governança de dados em aberto"),
            ("bullets", [
                "Registar URL/licença e versão do dataset original, se disponíveis.",
                "Guardar data de obtenção e resumo criptográfico do ficheiro de dados.",
                "Verificar a origem e validade das etiquetas, sobretudo em conflitos.",
                "Documentar consentimento, retenção e restrições de redistribuição."
            ])
        ]),
        ("5. Canvas: fontes, atributos e simulação", [
            ("p", "As características disponíveis provêm de duas realidades distintas: metadados estáticos do executável e variáveis comportamentais associadas ao dataset. Um ficheiro PE analisado sem execução permite obter estrutura do cabeçalho e secções, mas não revela diretamente chamadas de rede, alterações de registry ou comportamento futuro do processo."),
            ("p", "As características extraídas para um ficheiro podem, por isso, ser incompletas face às 72 colunas utilizadas no treino. O pipeline de preparação usa imputação/transformações ajustadas no treino. A imputação possibilita processamento, mas não converte um campo desconhecido numa observação real."),
            ("p", "A simulação económica não foi executada. Para não atribuir significado falso a números hipotéticos, este relatório não converte a matriz de confusão em euros, dólares ou incidentes evitados. O resultado só mostra contagens e métricas no teste."),
            ("table", (["Fonte", "Observável no protótipo", "Risco"], [
                ["Cabeçalho PE", "Sim, via leitura estática de PE", "Atributos legítimos e maliciosos podem sobrepor-se."],
                ["Secções .text/.rdata", "Alguns campos selecionados", "Secções podem faltar ou variar."],
                ["Comportamento em execução", "Não é recolhido pelo extrator estático", "Campos do dataset podem ficar ausentes/imputados."],
                ["Rótulo real do ficheiro novo", "Não conhecido no momento da previsão", "Sem rótulo não há medição contínua de desempenho."],
            ]))
        ]),
        ("6. Canvas: como se fazem previsões", [
            ("p", "O fluxo de utilização tem duas entradas possíveis. Num caso, recebe-se um CSV com as características esperadas pelo schema do modelo. Noutro, a ferramenta de extração lê um executável Windows ou um ZIP e produz uma linha de características PE estáticas para encaminhar para a previsão."),
            ("p", "O serviço HTTP local documentado expõe GET /health e POST /predict. O pedido de previsão aceita até 256 linhas JSON, com limite de 2 MiB, exige as características do schema e devolve resposta de erro explícita para JSON inválido, campos obrigatórios ausentes, tipo de conteúdo não JSON ou tamanho excedido. Por omissão, o servidor liga-se a 127.0.0.1."),
            ("p", "A classe/score deve ser interpretada como sinal para revisão, não como probabilidade calibrada de risco sem validação. Algumas implementações de modelos podem devolver score de decisão; comparar scores entre modelos sem calibração é inadequado."),
            ("subheading", "Exemplo operacional seguro"),
            ("p", "Um analista escolhe um executável suspeito, obtém características estáticas, consulta a previsão e encaminha o caso para revisão. Se o formato não for PE válido, o extrator deve devolver erro em vez de inventar uma classificação. O protótipo não executa a amostra nem verifica o seu comportamento em sandbox.")
        ]),
        ("7. Canvas: construção dos modelos", [
            ("p", "A tarefa supervisionada usa Class como alvo positivo Malware. Depois da limpeza, foram separados treino, validação e teste de forma estratificada, com semente 42 e proporções 60% / 20% / 20%: 8.331, 2.777 e 2.778 linhas, respetivamente."),
            ("p", "Doze classificadores foram comparados. A seleção inicial favoreceu recall Malware na validação, usando precisão e F1 como desempate. O teste ficou reservado. No modelo escolhido, a pesquisa de hiperparâmetros usou StratifiedKFold com três folds em treino + validação; o modelo não foi escolhido com base no resultado de teste."),
            ("p", "O melhor Random Forest reportado usa 100 árvores, profundidade máxima sem limite e seleção de 32 características. A validação cruzada obteve recall médio de Malware 0,9970 (desvio-padrão 0,0007) e precisão média 0,9862 (desvio-padrão 0,0020). O intervalo pequeno entre recall de treino e validação é um diagnóstico favorável, não garantia de generalização."),
            ("table", (["Fase", "Dados", "Uso permitido"], [
                ["Treino", "8.331 linhas", "Ajustar transformações e estimadores."],
                ["Validação", "2.777 linhas", "Comparar modelos e orientar seleção inicial."],
                ["CV da pesquisa", "Treino + validação", "Escolher parâmetros; teste continua isolado."],
                ["Teste", "2.778 linhas", "Avaliação final reportada uma vez."],
            ]))
        ]),
        ("8. Canvas: características e preparação", [
            ("p", "A EDA contabiliza 72 preditores: 63 numéricos e 9 categóricos. Identificadores criptográficos não são usados como preditores. Os campos numéricos inicialmente representados como texto, incluindo números hexadecimais, exigem conversão e validação; categorias têm codificação própria."),
            ("p", "Depois da limpeza, foram registados 485 valores ausentes. Para a classe binária restam 7.168 Benign e 6.718 Malware. A distribuição é suficientemente próxima para que a accuracy seja interpretável em conjunto com recall, precisão, F1 e métricas de ranking; ainda assim, não elimina viés de amostragem."),
            ("p", "As transformações aprendidas — imputação, codificação, seleção e eventual escala — devem ser ajustadas dentro do pipeline em cada partição/fold, e não antes de dividir os dados. A documentação metodológica do projeto descreve esse desenho. PCA aparece apenas como exploração: PC1+PC2 explicam 17,0% da variância e foram ajustados sobre os dados limpos; não é evidência de validação preditiva."),
            ("p", "Os testes univariados incluem 63 variáveis numéricas e 8 categóricas, com correção Benjamini–Hochberg. Associações e significância estatística são descritivas: não demonstram causalidade nem utilidade fora deste conjunto.")
        ]),
        ("9. Canvas: monitorização", [
            ("p", "Não foi demonstrado um sistema de monitorização em produção. Não há fluxo de rótulos posteriores, alertas de drift, painel operacional, SLA ou política de re-treino. A existência de um endpoint /health verifica disponibilidade do serviço/modelo, não a qualidade continuada das previsões."),
            ("subheading", "Plano mínimo antes de operação"),
            ("bullets", [
                "Guardar versão do modelo, schema, hash do artefacto e data de treino.",
                "Registar entradas e previsões com controlos de privacidade e retenção.",
                "Obter rótulos revistos por analistas e calcular recall, precisão e falsos negativos.",
                "Separar métricas por origem, família, período e formato de ficheiro.",
                "Monitorizar campos ausentes, valores fora de gama, distribuição de scores e taxa de erros.",
                "Definir critérios de suspensão, revisão e novo treino; não atualizar o modelo automaticamente."
            ]),
            ("p", "A frequência de revisão e os limiares de alerta dependem do volume e do custo operacional real. Não foram fixados neste relatório porque não existem medições de produção que os justifiquem.")
        ]),
        ("10. EDA: o que os dados mostram", [
            ("p", "A exploração confirma que os dados não são simplesmente a tabela original sem alterações: 21.752 linhas passam a 13.886 após deduplicação e remoção de 611 registos com conflito de etiqueta. A composição final contém 7.168 Benign, 2.253 RAT, 1.795 Stealer, 1.476 Ransomware e 1.194 Trojan; as últimas categorias somam Malware."),
            ("figure", ("results/roadmap/eda/class_distribution.png", "Figura 1 — Distribuição das duas classes depois da limpeza.", 5.5)),
            ("p", "A proximidade entre as duas barras é coerente com a contagem final quase equilibrada. Não significa que cada família esteja igualmente representada: por exemplo, a documentação EDA lista apenas 21 amostras da família Ragnar, enquanto Snake tem 521. A avaliação binária pode ocultar desempenho fraco em famílias pequenas."),
            ("p", "Foram calculados 1.953 pares de correlação de Pearson e 1.953 de Spearman para 63 variáveis numéricas. A EDA é útil para procurar redundância, outliers e associações, mas não substitui validação por grupos nem validação temporal.")
        ]),
        ("11. Preparação dos dados e reprodutibilidade", [
            ("p", "A limpeza documentada verifica campos requeridos, remove duplicados por MD5 e lida com rótulos conflitantes. A tabela final tem uma observação por hash, reduzindo o risco de a mesma amostra aparecer em treino e teste. Esta proteção é necessária, mas não prova independência entre amostras relacionadas ou variantes próximas."),
            ("p", "Um artefacto de preparação separado relata um split 80/20 para análises próprias. Esse split não deve ser confundido com o benchmark supervisionado, que utiliza 60/20/20. Os resultados apresentados nas páginas seguintes referem-se exclusivamente ao protocolo 60/20/20 do relatório supervisionado."),
            ("p", "A aleatoriedade controlada (random_state=42) facilita repetir a partição. A repetibilidade integral também depende de versão de bibliotecas, origem do CSV, código e artefacto de modelo. O relatório descreve a configuração encontrada nos artefactos atuais; não afirma que todos estes elementos estejam arquivados num pacote imutável."),
            ("table", (["Controlo", "Estado / valor"], [
                ["Linhas antes / depois", "21.752 / 13.886"],
                ["Conflitos de etiqueta removidos", "611"],
                ["Features / valores ausentes", "72 / 485"],
                ["Partição do benchmark", "8.331 / 2.777 / 2.778"],
                ["PCA de duas componentes", "17,0% de variância; apenas descritivo"],
            ]))
        ]),
        ("12. Comparação de classificadores", [
            ("p", "Na validação, o Random Forest liderou no critério definido: recall de Malware. As métricas seguintes pertencem à comparação de validação, não ao teste final. A precisão-recall e ROC-AUC elevadas devem ser lidas no contexto do mesmo dataset, da mesma preparação e do mesmo desenho de partição."),
            ("table", (["Modelo", "Recall Malware", "Precisão Malware", "F1 Malware"], [
                ["Random Forest", "0,9993", "0,9904", "0,9948"],
                ["Linear SVM", "0,9985", "0,9803", "0,9893"],
                ["Gradient Boosting", "0,9978", "0,9831", "0,9904"],
                ["Extra Trees", "0,9978", "0,9831", "0,9904"],
                ["AdaBoost", "0,9978", "0,9817", "0,9897"],
            ])),
            ("p", "A tabela é uma seleção dos cinco modelos de maior desempenho relevante no ficheiro de comparação. Modelos diferentes podem reportar scores de decisão ou probabilidades; o relatório identifica o tipo de score. Um score de decisão de SVM não deve ser interpretado diretamente como probabilidade."),
            ("p", "A seleção prioriza reduzir falsos negativos, mas tem custo em falsos positivos. A escolha operacional requer prevalência, capacidade dos analistas e custo de erro, que não foram medidos.")
        ]),
        ("13. Avaliação final no teste", [
            ("p", "O conjunto de teste tem 2.778 linhas (1.434 Benign e 1.344 Malware). Reproduzi a previsão com o artefacto Random Forest e a mesma divisão estratificada indicada no relatório; a matriz observada é a seguinte."),
            ("figure", ("results/roadmap/model_evaluation/supervised_confusion_matrix.png", "Figura 2 — Matriz de confusão do teste; linhas = classe real, colunas = classe prevista.", 4.65)),
            ("table", (["Métrica (classe positiva: Malware)", "Resultado"], [
                ["Accuracy", "0,9924"],
                ["Precisão", "0,9875"],
                ["Recall", "0,9970"],
                ["F1", "0,9922"],
                ["ROC-AUC / PR-AUC", "0,9990 / 0,9989"],
                ["Matriz: TN / FP / FN / TP", "1.417 / 17 / 4 / 1.340"],
            ])),
            ("p", "O recall corresponde a 1.340 de 1.344 casos Malware reconhecidos; os quatro restantes são falsos negativos neste teste. Os 17 falsos positivos são amostras Benign sinalizadas como Malware. Nem estas contagens nem a taxa observada antecipam necessariamente o desempenho em tráfego futuro.")
        ]),
        ("14. Curva de aprendizagem e capacidade de ranking", [
            ("figure", ("results/roadmap/model_evaluation/learning_curve.png", "Figura 3 — Recall de treino e validação ao variar o número de exemplos por fold.", 5.45)),
            ("p", "A curva reporta recall de treino próximo de 1,0 e recall de validação próximo de 0,996–0,997 em todos os tamanhos avaliados. O pequeno intervalo é consistente com o diagnóstico reportado de ausência de sinal forte de underfitting ou overfitting segundo essa métrica. A curva usa treino + validação, não o teste."),
            ("figure", ("results/roadmap/model_evaluation/supervised_roc_curve.png", "Figura 4 — Curva ROC do Random Forest; área aproximada de 0,999.", 4.5)),
            ("p", "ROC-AUC resume a ordenação entre positivos e negativos para diferentes limiares; não fixa um limiar de decisão nem informa diretamente quantos alertas uma equipa pode tratar. Para escolher limiar, medir curvas de precisão-recall, custos e capacidade de revisão em dados representativos.")
        ]),
        ("15. Exemplo concreto: ficheiro PE", [
            ("p", "O extrator percorre um caminho estático: valida o cabeçalho MZ, analisa estrutura PE e recolhe propriedades selecionadas dos cabeçalhos DOS/PE e secções .text/.rdata. Para um ficheiro válido pode produzir atributos como EntryPoint, arquitetura, tamanho de imagem e dimensões de secção; em caso de formato inválido devolve erro."),
            ("p", "Os limites implementados incluem 50 MiB por PE, 250 MiB por arquivo ZIP, 50 MiB por membro, 500 MiB de tamanho total descomprimido e 500 membros. Estes limites restringem consumo e exposição a arquivos excessivos. O extrator não executa o ficheiro nem constitui uma sandbox."),
            ("p", "Exemplo de interpretação: se a saída for Malware, o resultado indica semelhança estatística com exemplos rotulados Malware durante treino. O passo correto é rever amostra, score e contexto. Não se conclui que é ransomware, que está ativo ou que um resultado Benign é seguro."),
            ("subheading", "Desfasamento de características"),
            ("p", "O dataset contém campos de comportamento que só seriam conhecidos através de telemetria ou execução controlada. A extração estática não os preenche com observações dinâmicas. Quando o pipeline imputa valores, isso pode alterar o comportamento do modelo; é necessário medir o desempenho no modo real de entrada, não apenas com linhas completas do dataset.")
        ]),
        ("16. Arquitetura de utilização", [
            ("p", "A aplicação Tkinter fornece uma interface Windows para escolher CSV, instalar dependências, treinar e avaliar modelos, fazer previsões, extrair características de .exe/.zip, iniciar a API e controlar o bot Telegram. O módulo de serviço disponibiliza uma API HTTP local e o módulo de previsão carrega o pipeline e schema do modelo."),
            ("table", (["Componente", "Função", "Estado que se pode afirmar"], [
                ["Preparação / treino", "Transformação, comparação e avaliação supervisionada.", "Implementado e com artefactos locais."],
                ["Extrator PE", "Produzir atributos estáticos de executáveis / ZIP.", "Implementado; não executa ficheiros."],
                ["API", "Health check e previsões JSON.", "Documentada para uso local."],
                ["Aplicação desktop", "Interface para fluxos de treino e inferência.", "Existe no código do projeto."],
                ["Bot Telegram", "Interface conversacional opcional.", "Disponível; depende de token/configuração externa."],
            ])),
            ("p", "O bot conversacional pode usar um modelo de linguagem local/remoto configurado, distinto do classificador Random Forest. Uma resposta textual desse modelo não é uma previsão de segurança e não deve ser confundida com o benchmark binário."),
            ("p", "O servidor de inferência liga a 127.0.0.1 por omissão. A documentação assinala corretamente que expor o protótipo à rede exigiria autenticação, TLS e controlos operacionais ainda não implementados.")
        ]),
        ("17. Validade, limitações e riscos", [
            ("table", (["Risco / limite", "Implicação", "Mitigação necessária"], [
                ["Proveniência e etiquetas não verificadas independentemente", "Viés ou erros podem contaminar treino/teste.", "Arquivar fonte, versões e regras de anotação."],
                ["Divisão aleatória", "Pode não simular famílias ou períodos futuros.", "Teste temporal e divisão por família/grupo."],
                ["Amostras relacionadas", "Hashes únicos não garantem independência familiar.", "Deduplicar por hash e agrupar por origem/linhagem."],
                ["Campos dinâmicos ausentes na extração", "Inferência real pode diferir do benchmark.", "Avaliar entradas produzidas pelo próprio extrator."],
                ["Classes de família desequilibradas", "Métrica binária mascara famílias raras.", "Relatar recall por família e intervalo de incerteza."],
                ["Artefacto joblib", "Modelo serializado não confiável pode executar código.", "Carregar apenas artefactos gerados localmente."],
            ])),
            ("p", "O dataset contém 26 famílias de malware segundo a EDA; algumas têm poucas dezenas de exemplos. Não se afirma desempenho por família, nem há teste externo independente. A remoção de duplicados por MD5 é útil, mas amostras da mesma família podem partilhar padrões e continuar distribuídas entre partições."),
            ("p", "As métricas altas são um resultado real do protocolo, mas a facilidade do conjunto, vazamento residual, origem ou similaridade entre registos exigem validação adicional antes de qualquer conclusão sobre generalização.")
        ]),
        ("18. Ética, segurança e uso responsável", [
            ("p", "Uma classificação errada pode bloquear software legítimo ou deixar passar malware. O protótipo deve ser apresentado como ferramenta de apoio experimental. As previsões precisam de revisão por pessoa competente, especialmente antes de ações irreversíveis."),
            ("p", "Ficheiros suspeitos devem ser tratados como potencialmente perigosos. A análise estática não executa o ficheiro, mas isso não elimina riscos de parsers, bibliotecas, arquivos malformados ou gestão indevida de amostras. O projeto aplica limites de tamanho; qualquer análise dinâmica exigiria isolamento apropriado."),
            ("p", "Hashes e amostras podem ser dados sensíveis para organizações. Uma implementação futura deve definir acesso, retenção, logging, partilha e procedimento de incidentes. As credenciais do bot Telegram devem permanecer em variáveis de ambiente e o acesso deve ser restrito por IDs autorizados; não se devem colocar tokens em código ou no relatório."),
            ("p", "A explicabilidade atual é limitada: a classificação binária não fornece, por si só, uma razão causal. Importância de atributos ou explicações locais, caso sejam acrescentadas, devem ser validadas e não descritas como prova de intenção maliciosa.")
        ]),
        ("19. Próximos passos recomendados", [
            ("table", (["Prioridade", "Trabalho", "Critério de conclusão"], [
                ["1 — Dados", "Confirmar fonte/licença, datas e etiquetas; versionar dataset.", "Proveniência auditável e política de retenção."],
                ["2 — Generalização", "Separar por hash e família, testar por período e em dataset externo.", "Relatório independente com intervalos por grupo."],
                ["3 — Modo real", "Avaliar ficheiros convertidos pelo extrator, incluindo ausências.", "Métricas e taxa de erro para entradas reais."],
                ["4 — Decisão", "Escolher limiar com analistas e custos observados.", "Política documentada; aprovação humana mantida."],
                ["5 — Operação", "Autenticação/TLS, controlo de acesso, auditoria e monitorização.", "Checklist de segurança e ensaio de operação."],
                ["6 — Multiclasse", "Explorar categoria/família com suporte mínimo definido.", "Resultados por classe e limitações declaradas."],
            ])),
            ("p", "Uma sequência prudente é resolver primeiro os riscos de dados e validação, antes de otimizar o modelo ou automatizar decisões. Aumentar recall no mesmo split não substitui evidência independente.")
        ]),
        ("20. Conclusões", [
            ("p", "O trabalho integra preparação, EDA, comparação supervisionada e um protótipo de utilização para classificação Benign/Malware. O processo documenta a remoção de conflitos e duplicados, separa um teste reservado e dá prioridade ao recall de Malware."),
            ("p", "No teste do protocolo descrito, o Random Forest obteve 0,9970 de recall, 0,9875 de precisão e 0,9922 de F1 para Malware, com 4 falsos negativos e 17 falsos positivos em 2.778 linhas. A reprodução do modelo e da partição confirmou a matriz indicada. Estes resultados são promissores como prova de conceito, mas não são uma certificação."),
            ("p", "A designação do projeto não deve levar a concluir que o modelo reconhece exclusivamente ransomware: a etiqueta positiva inclui RAT, Stealer e Trojan. Também não há prova de eficácia em ficheiros atuais, famílias inéditas ou telemetria de uma organização."),
            ("p", "A conclusão prática é manter o sistema em contexto de investigação e triagem assistida, fechar a proveniência dos dados e executar validação por família/tempo e entradas reais. Só depois dessas etapas se deve discutir um uso operacional.")
        ]),
        ("21. Transparência, referências e verificação final", [
            ("subheading", "Declaração de apoio de IA"),
            ("p", "Foi utilizado um assistente de IA para organizar a estrutura, redigir e rever partes deste relatório e para ajudar a sintetizar artefactos do projeto. Os valores quantitativos foram confrontados com os relatórios locais; a matriz de confusão foi reproduzida com o modelo e a partição documentados. A equipa deve rever o texto, confirmar autoria e contexto académico, preencher a identificação da capa e assumir responsabilidade pela versão entregue."),
            ("subheading", "Fontes consultadas no projeto"),
            ("bullets", [
                "docs/business-understanding.md — objetivos, utilizadores, critérios e riscos.",
                "docs/data-understanding.md — descrição do dataset e atributos.",
                "docs/deployment.md — API local, limites e segurança operacional.",
                "results/roadmap/eda/eda_summary.json — limpeza, distribuição, testes e PCA.",
                "results/roadmap/model_evaluation/supervised_learning_report.json — desenho experimental e métricas.",
                "results/roadmap/model_evaluation/supervised_model_comparison.csv — comparação de validação.",
                "src/extract_features.py, src/app.py e src/serve.py — implementação do protótipo."
            ]),
            ("p", "Estrutura de enquadramento: Louis Dorard, Machine Learning Canvas v1.2, OWNML, 2015, licença Creative Commons Attribution-ShareAlike 4.0 International (CC BY-SA 4.0). O relatório segue as áreas do Canvas sem reproduzir o formulário integral. Referência: https://creativecommons.org/licenses/by-sa/4.0/"),
            ("p", "Antes da entrega: confirmar nomes, curso, unidade curricular, docente e data; validar o relatório contra os requisitos institucionais; rever a declaração de IA e as referências; abrir o DOCX e verificar que a paginação final permanece abaixo do limite solicitado.")
        ]),
    ]

    for title, contents in pages:
        add_page(document, title, contents)

    document.core_properties.title = "Deteção de malware com aprendizagem automática"
    document.core_properties.subject = "Relatório técnico e de enquadramento do projeto IAAC"
    document.core_properties.author = "Equipa do projeto (a preencher)"
    document.core_properties.keywords = "malware, ransomware, aprendizagem automática, relatório"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_report())
